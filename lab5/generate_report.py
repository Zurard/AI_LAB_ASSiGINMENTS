"""Build the Lab 5 Word report from a live GMM-EM and K-means run."""

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
import numpy as np

from experiment import NAMES, SPECS, run, save_figures, summarize


HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Lab_Assignment_5_GMM_EM_and_KMeans.docx")


def set_run(run, name="Calibri", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8FA4B8")
        borders.append(el)
    tcPr.append(borders)


def para(doc, text, *, size=12, bold=False, center=False, space_after=8, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first_line:
            p.paragraph_format.first_line_indent = Cm(0.75)
    set_run(p.add_run(text), size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        run.font.name = "Calibri"
    return h


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    set_run(p.add_run(text), size=10, italic=True, color=(80, 80, 80))


def add_picture(doc, path, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(path, width=Inches(width))


def add_code(doc, code, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    set_run(p.add_run(title), size=11, bold=True, italic=True, color=(31, 58, 95))
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F6F8")
    set_cell_border(cell)
    cell.text = ""
    cp = cell.paragraphs[0]
    cp.paragraph_format.first_line_indent = Cm(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.05
    set_run(cp.add_run(code.rstrip() + "\n"), name="Consolas", size=8, color=(30, 30, 30))


def fill_header_row(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = ""
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cell.paragraphs[0].add_run(text), size=10, bold=True, color=(255, 255, 255))
        shade_cell(cell, "1F3A5F")
        set_cell_border(cell)


def fill_row(row, texts, center=True, shade=None):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = ""
        if center:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cell.paragraphs[0].add_run(str(text)), size=10)
        if shade:
            shade_cell(cell, shade)
        set_cell_border(cell)


def snippet(path, start, end):
    lines = open(path, encoding="utf-8").read().splitlines()
    return "\n".join(lines[start - 1:end])


def build():
    print("Running GMM-EM and K-means...")
    data = save_figures(run())
    fig = data["fig_dir"]
    gmm = data["gmm"]
    rows = summarize(data)
    means = data["means_aligned"]
    stds = np.sqrt(data["variances_aligned"])
    weights = data["weights_aligned"]
    n_dis = int(data["disagree"].sum())

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(12)

    banner = doc.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(banner.add_run("ARTIFICIAL INTELLIGENCE LABORATORY"),
            size=13, bold=True, color=(31, 58, 95))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    set_run(title.add_run(
        'Lab Assignment 5 : "Clustering with Gaussian Mixture EM and K-means"'
    ), size=16, bold=True, color=(31, 58, 95))

    meta = [
        ("Name", "Shreetej Meshram"),
        ("ID", "202352333"),
        ("Topic", "Unsupervised clustering"),
        ("Methods", "GMM-EM (diagonal covariance) and K-means"),
        ("Data", "3 Gaussian clusters, 150 points each, 2 dimensions"),
        ("Language", "Python 3 (NumPy and Matplotlib only)"),
    ]
    t = doc.add_table(rows=len(meta), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        fill_row(t.rows[i], [k, v], center=False, shade="F4F6F8" if i % 2 == 0 else None)
        t.rows[i].cells[0].paragraphs[0].runs[0].bold = True
    para(doc, "", first_line=False, space_after=4)

    heading(doc, "1. Aim", 1)
    para(doc,
         "The aim is to cluster an unlabeled two-dimensional point set with a "
         "Gaussian mixture model trained by Expectation Maximization, and to "
         "compare that clustering with K-means on the same points. The data "
         "are three Gaussian clusters with different means and different "
         "standard deviations, 150 points in each cluster. The mixture uses "
         "diagonal covariance matrices, so each cluster may be stretched along "
         "x or along y, but not rotated. The comparison is used to show where "
         "the two methods disagree, and why the shape and the variance of a "
         "cluster cause that disagreement.")

    heading(doc, "2. Question Discussion", 1)
    heading(doc, "2.1 The data", 2)
    para(doc,
         "Each cluster is drawn independently from a normal distribution. "
         "The first is compact and nearly round. The second is wide along x "
         "and thin along y. The third is thin along x and tall along y. "
         "Because the covariance is diagonal, the level curves of each "
         "density are axis-aligned ellipses. Table 1 is the generator. "
         "The seed is 7, so the same 450 points are used for every figure "
         "in this report. After the points are drawn they are shuffled, and "
         "the true labels are kept only for the plots and the accuracy "
         "check. Both algorithms see the coordinates alone.")

    heading(doc, "2.2 GMM-EM with a diagonal covariance", 2)
    para(doc,
         "A Gaussian mixture says each point was produced by first picking "
         "a cluster with probability pi_k and then drawing from that "
         "cluster's Gaussian. With a diagonal covariance the density of "
         "cluster k at a point x is the product of two independent "
         "one-dimensional normals, one for each coordinate. The hidden "
         "variable is the cluster index. EM treats it as a fractional count.")
    para(doc,
         "Initialization picks three distinct data points as the means, "
         "sets every mixing weight to 1/3, and sets every cluster's "
         "variance vector to the variance of the whole cloud. The E-step "
         "computes the responsibility of each cluster for each point: the "
         "normalized product of the weight and the Gaussian density. The "
         "M-step updates the weight to the average responsibility, the mean "
         "to the responsibility-weighted average of the points, and each "
         "diagonal variance to the weighted mean squared deviation. The "
         "loop stops when the change in log-likelihood falls below 1e-5, "
         "or after 80 iterations. Eight random starts are run and the start "
         "with the highest final log-likelihood is kept, because EM can "
         "stop in a poor local maximum.")

    heading(doc, "2.3 K-means, and why it is a different model", 2)
    para(doc,
         "K-means alternates two steps as well. Each point is given to the "
         "nearest center under Euclidean distance, then each center is moved "
         "to the mean of its points. Fifteen random starts are used and the "
         "run with the smallest sum of squared distances is kept. Euclidean "
         "distance treats every direction the same and does not have a "
         "per-cluster variance. The decision boundary between two centers is "
         "the perpendicular bisector of the segment that joins them. A "
         "Gaussian with a large variance along x will still claim a point "
         "that sits far to the side of its mean, because that deviation is "
         "cheap under the Mahalanobis distance (x - mu)^2 / sigma^2. K-means "
         "will give that same point to whichever center is closer in ordinary "
         "distance, often the compact cluster. That is the disagreement this "
         "lab is built to show.")

    heading(doc, "3. Code", 1)
    para(doc,
         "gmm.py is the mixture. kmeans.py is Lloyd's algorithm. "
         "experiment.py draws the data, fits both models, aligns the "
         "arbitrary cluster indices to the true labels, and writes the "
         "figures. The snippets are the steps the assignment asks for: "
         "initialization, the E-step, the M-step, the convergence test, "
         "and the K-means update.")

    heading(doc, "3.1 Initialization, E-step, M-step, convergence", 2)
    add_code(doc, snippet(os.path.join(HERE, "gmm.py"), 24, 54),
             "Snippet 1  -  gmm.py, one EM run")
    para(doc,
         "log_gaussian_diag evaluates the diagonal normal in the log domain, "
         "so a very small density does not underflow to zero before the "
         "clusters are compared. Subtracting the row maximum before the "
         "exponential is the stable way to turn those logs into "
         "responsibilities. The convergence test is the absolute change in "
         "the observed-data log-likelihood. Variances are floored at 1e-8 "
         "so a cluster cannot collapse onto a single point and make the "
         "density infinite.")

    heading(doc, "3.2 K-means", 2)
    add_code(doc, snippet(os.path.join(HERE, "kmeans.py"), 14, 32),
             "Snippet 2  -  kmeans.py, assignment and center update")
    para(doc,
         "The squared Euclidean distance is used only to choose the nearest "
         "center. The new center of a cluster is the coordinate-wise mean of "
         "the points assigned to it. If a start produces an empty cluster, "
         "that center is left where it was. The inertia of a finished run is "
         "the sum of squared distances to the assigned centers, and the "
         "smallest inertia across starts is the K-means solution reported here.")

    heading(doc, "4. Results", 1)
    heading(doc, "4.1 True clusters and the unlabeled cloud", 2)
    para(doc,
         "Figure 1 shows the same 450 points twice. On the left they are "
         "coloured by the cluster that generated them. On the right the "
         "labels are hidden. The wide cluster stretches across the lower "
         "right. The tall cluster stretches up the left side. The compact "
         "cluster sits near the origin, in the corner where those two arms "
         "would meet if they were longer. That corner is where a round "
         "distance and an elliptical density will not agree.")
    add_picture(doc, os.path.join(fig, "true_and_unlabeled.png"), 6.4)
    caption(doc, "Figure 1. True clusters (left) and the unlabeled dataset given to both algorithms (right).")

    spec_table = doc.add_table(rows=4, cols=4)
    spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fill_header_row(spec_table.rows[0], ["Cluster", "Mean (x1, x2)", "Std (x1, x2)", "Points"])
    for i, (spec, name) in enumerate(zip(SPECS, NAMES)):
        fill_row(
            spec_table.rows[i + 1],
            [name,
             "(%.2f, %.2f)" % tuple(spec["mean"]),
             "(%.2f, %.2f)" % tuple(spec["std"]),
             "150"],
            shade="F4F6F8" if i % 2 == 0 else None,
        )
    caption(doc, "Table 1. Generating distributions. Seed 7. Covariances are diagonal.")

    heading(doc, "4.2 Fitted mixture and the log-likelihood", 2)
    para(doc,
         "The winning EM run used %d iterations. The log-likelihood of the "
         "450 points moved from %.2f after the first E-step to %.2f at "
         "convergence. Figure 3 is that curve: it rises and then flattens, "
         "which is what the convergence test is waiting for. Figure 2 shows "
         "the hard assignments (each point given to its largest "
         "responsibility) together with the 1-standard-deviation and "
         "2-standard-deviation ellipses of the fitted diagonal Gaussians. "
         "The ellipses recover the shape of the generator. Table 2 puts the "
         "fitted parameters next to the values that were used to draw the data."
         % (len(gmm.loglik_), gmm.loglik_[0], gmm.loglik_[-1]))
    add_picture(doc, os.path.join(fig, "gmm_fit.png"), 5.5)
    caption(doc, "Figure 2. GMM-EM hard assignments. Ellipses are the fitted 1-sigma and 2-sigma contours. Crosses are the fitted means.")
    add_picture(doc, os.path.join(fig, "loglik.png"), 5.8)
    caption(doc, "Figure 3. Observed-data log-likelihood of the best EM start.")

    ptab = doc.add_table(rows=4, cols=5)
    ptab.alignment = WD_TABLE_ALIGNMENT.CENTER
    fill_header_row(ptab.rows[0], ["Cluster", "True mean", "Fitted mean", "True std", "Fitted std"])
    for i in range(3):
        fill_row(ptab.rows[i + 1], [
            NAMES[i],
            "(%.2f, %.2f)" % tuple(SPECS[i]["mean"]),
            "(%.2f, %.2f)" % tuple(means[i]),
            "(%.2f, %.2f)" % tuple(SPECS[i]["std"]),
            "(%.2f, %.2f)" % tuple(stds[i]),
        ], shade="F4F6F8" if i % 2 == 0 else None)
    caption(doc, "Table 2. True parameters and the EM estimate. Component indices are aligned to the true labels.")

    para(doc,
         "After alignment, the hard GMM labeling matches the true cluster "
         "on %.1f%% of the points. The fitted standard deviations stay close "
         "to the generator: the wide cluster is still wide in x (fitted "
         "std %.2f against true %.2f) and the tall cluster is still tall in "
         "y (fitted std %.2f against true %.2f). The mixing weights are "
         "%.2f, %.2f and %.2f, close to the equal 1/3 that was used to draw "
         "the sample."
         % (100 * data["gmm_acc"], stds[1, 0], SPECS[1]["std"][0],
            stds[2, 1], SPECS[2]["std"][1],
            weights[0], weights[1], weights[2]))

    heading(doc, "4.3 K-means on the same points", 2)
    para(doc,
         "Figure 4 is the K-means partition. The crosses are the three "
         "centers. The inertia of the chosen start is %.1f. Matched against "
         "the true labels, K-means is correct on %.1f%% of the points, "
         "against %.1f%% for the mixture. The missing points are not "
         "scattered at random. They lie on the long arms."
         % (data["km"]["inertia"], 100 * data["km_acc"], 100 * data["gmm_acc"]))
    add_picture(doc, os.path.join(fig, "kmeans.png"), 5.5)
    caption(doc, "Figure 4. K-means assignments on the unlabeled data. Crosses are the learned centers.")

    heading(doc, "4.4 Where the two assignments differ", 2)
    para(doc,
         "After both labelings are aligned to the true clusters, they "
         "disagree on %d of the 450 points. Figure 5 marks those points in "
         "the colour of the cluster that actually generated them, and leaves "
         "the agreeing points grey. Table 3 counts the disagreements inside "
         "each true cluster. Almost all of them belong to the two elongated "
         "clusters. The compact cluster contributes only %d points."
         % (n_dis, rows[0]["disagree"]))
    add_picture(doc, os.path.join(fig, "disagree.png"), 5.5)
    caption(doc, "Figure 5. Points whose K-means label and GMM-EM label differ. Colour is the true cluster.")

    dtab = doc.add_table(rows=5, cols=3)
    dtab.alignment = WD_TABLE_ALIGNMENT.CENTER
    fill_header_row(dtab.rows[0], ["True cluster", "Points", "K-means differs from GMM"])
    for i, row in enumerate(rows):
        fill_row(dtab.rows[i + 1], [row["name"], str(row["n"]), str(row["disagree"])],
                 shade="F4F6F8" if i % 2 == 0 else None)
    fill_row(dtab.rows[4], ["total", "450", str(n_dis)], shade="FFF3CD")
    caption(doc, "Table 3. Disagreements by the cluster that generated the point.")

    para(doc,
         "The reason is the shape. A point on the far left of the wide "
         "cluster can be closer, in Euclidean distance, to the compact "
         "center near the origin than to the center of the wide cluster. "
         "K-means therefore assigns it to the compact cluster. The mixture "
         "does not. Its fitted standard deviation along x for that cluster "
         "is about %.2f, so a large horizontal deviation is still likely "
         "under that Gaussian, while the same deviation is unlikely under "
         "the compact Gaussian, whose standard deviation is about %.2f in "
         "both directions. The same thing happens on the vertical arm: a "
         "point high above the tall cluster is far from every center in "
         "raw distance, but it is cheap for the tall Gaussian because that "
         "Gaussian's variance along y is large. K-means has no variance "
         "parameter, so it cannot represent that stretch. It cuts the cloud "
         "with straight bisectors and clips the tails."
         % (stds[1, 0], stds[0, 0]))
    para(doc,
         "The two methods agree on the dense cores, where a point is close "
         "to its own mean in every direction. They separate exactly where "
         "the clusters stop being round. That is the content of Table 3: "
         "%d disagreements on the wide cluster, %d on the tall cluster, and "
         "%d on the compact cluster."
         % (rows[1]["disagree"], rows[2]["disagree"], rows[0]["disagree"]))

    heading(doc, "5. Result", 1)
    para(doc,
         "Three two-dimensional Gaussian clusters were drawn, 150 points "
         "each, with different means and different axis-aligned standard "
         "deviations. GMM-EM with diagonal covariances recovered those "
         "shapes: the log-likelihood rose from %.2f to %.2f in %d iterations, "
         "and the hard labeling matched the true clusters on %.1f%% of the "
         "points. K-means on the same unlabeled set matched the truth on "
         "%.1f%%. The two labelings differ on %d points, and %d of those "
         "points belong to the two elongated clusters. K-means fails there "
         "because Euclidean distance ignores variance. The mixture keeps "
         "those tails because a diagonal covariance charges a deviation "
         "only relative to that cluster's own standard deviation."
         % (gmm.loglik_[0], gmm.loglik_[-1], len(gmm.loglik_),
            100 * data["gmm_acc"], 100 * data["km_acc"],
            n_dis, n_dis - rows[0]["disagree"]))

    heading(doc, "5.1 Conclusion", 2)
    para(doc,
         "EM is the right fit for a Gaussian mixture because the cluster "
         "index is hidden: the E-step estimates it, and the M-step updates "
         "the weight, the mean and the diagonal variance. K-means is the "
         "special case in which every cluster is forced to be spherical and "
         "of equal spread, and the assignment is hard at every step. On "
         "this sample the cores of the three clusters look the same to both "
         "methods. The arms do not. The plots and Table 3 are the evidence: "
         "the disagreements sit on the high-variance axes, which is where a "
         "diagonal Gaussian and a Euclidean center stop describing the same "
         "set of points.")

    heading(doc, "5.2 Files submitted", 2)
    files = [
        ("gmm.py", "Diagonal GMM: initialization, E-step, M-step, convergence."),
        ("kmeans.py", "Lloyd's K-means with several random starts."),
        ("experiment.py", "Data, both fits, disagreement counts, figures."),
        ("main.py", "Prints the numerical results."),
        ("generate_report.py", "Builds this document from a live run."),
        ("requirements.txt", "numpy, matplotlib, python-docx."),
        ("figures/", "True data, GMM fit, log-likelihood, K-means, disagreements."),
        ("Lab_Assignment_5_GMM_EM_and_KMeans.docx", "This report."),
    ]
    ftab = doc.add_table(rows=1 + len(files), cols=2)
    ftab.alignment = WD_TABLE_ALIGNMENT.CENTER
    fill_header_row(ftab.rows[0], ["File", "Role"])
    for i, (name, role) in enumerate(files):
        fill_row(ftab.rows[i + 1], [name, role], center=False,
                 shade="F4F6F8" if i % 2 == 0 else None)

    doc.save(OUT)
    print("Wrote", OUT)
    print("disagree", n_dis, "gmm", round(data["gmm_acc"], 3), "km", round(data["km_acc"], 3))


if __name__ == "__main__":
    build()
